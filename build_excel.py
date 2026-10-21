"""Build subscription_forecast_excel_model.xlsx: a formula-driven Excel version
of the cohort forecast. Change any blue input and the whole model recalculates.

Driver values come from the Python model fitted on Jan 2022 - Dec 2025 history;
the Excel FY totals are checked against the Python forecast on the Summary tab.
"""
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from subrev import analysis as A
from subrev.data import N_MONTHS, fx_table, simulate

H = 24                                   # forecast months: Jan 2026 .. Dec 2027
subs = simulate()
fx = fx_table(N_MONTHS + H)
fc, p = A.forecast(subs, fx, fit_end=N_MONTHS - 1, horizon=H)
cmix = A.currency_mix(subs, fx, N_MONTHS - 1)
non_usd = 1 - float(cmix.get("USD", 0.0))
py_fy26 = fc[fc.month.str.startswith("2026")].fc_net_revenue_usd.sum()
py_fy27 = fc[fc.month.str.startswith("2027")].fc_net_revenue_usd.sum()

F = "Arial"
BLUE, BLACK, GREEN = Font(name=F, color="0000FF"), Font(name=F), Font(name=F, color="008000")
BOLD, TITLE = Font(name=F, bold=True), Font(name=F, bold=True, size=14)
HDR = Font(name=F, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3864")
YELLOW = PatternFill("solid", fgColor="FFFF00")
thin = Border(top=Side(style="thin"))
USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
SRC = "Fitted by the Python cohort model (run.py) on simulated Jan 2022 - Dec 2025 history."

wb = Workbook()

# ---------------- Inputs ---------------------------------------------------
ws = wb.active
ws.title = "Inputs"
ws["A1"], ws["A1"].font = "Subscription Revenue Forecast - Driver Inputs", TITLE
ws["A2"] = "Blue = input you can change. Yellow = key assumption. Black = formula. Green = link to another sheet."
ws["A2"].font = Font(name=F, italic=True)
inputs = [  # key, label, value, fmt, key_assumption, note
    ("trials", "Trial starts, Jan 2026 (trend, before seasonality)", p["signups_trend_next"], NUM, False, SRC),
    ("growth", "Trial starts growth, month over month", p["signup_growth_mom"], PCT, True, SRC),
    ("dec", "December seasonality factor", p["dec_factor"], "0.00x", False, SRC),
    ("dec25", "Trial starts in Dec 2025 (convert in Jan 2026)", p["trials_in_fit_end"], NUM, False, "Actual Dec 2025 trial starts."),
    ("conv", "Trial-to-paid conversion", p["trial_conversion"], PCT, True, SRC + " Trailing 6 months."),
    ("ann_share", "Share of new payers choosing annual plan", p["annual_share"], PCT, False, SRC),
    ("haz_a", "Churn curve a  (churn = a x tenure^b)", p["hazard_a"], "0.000", True, SRC),
    ("haz_b", "Churn curve b", p["hazard_b"], "0.000", True, SRC),
    ("renew", "Annual-plan renewal rate", p["annual_renewal"], PCT, True, SRC),
    ("arpu_m", "Net ARPU, monthly plan (USD / month, after fees)", p["arpu_monthly_usd"], USD2, False, SRC + " Dec 2025 FX."),
    ("arpu_a", "Net ARPU, annual plan (USD / month recognized)", p["arpu_annual_usd"], USD2, False, SRC + " Dec 2025 FX."),
    ("non_usd", "Non-USD share of revenue", non_usd, PCT, False, SRC + " Dec 2025 currency mix."),
    ("scenario", "FX scenario (pick from list)", "Base", None, True, "Choose Base, USD +10% or USD -10%."),
]
ref = {}
for i, (k, label, val, fmt, key, note) in enumerate(inputs, start=4):
    ws.cell(i, 1, label).font = BLACK
    c = ws.cell(i, 2, val)
    c.font = BLUE
    if fmt:
        c.number_format = fmt
    if key:
        c.fill = YELLOW
    c.comment = Comment(note, "Model")
    ref[k] = f"Inputs!$B${i}"
r0 = 4 + len(inputs) + 1
ws.cell(r0, 1, "FX scenario table").font = BOLD
ws.cell(r0 + 1, 1, "Scenario").font = BOLD
ws.cell(r0 + 1, 2, "Non-USD currencies vs USD").font = BOLD
for j, (name, shock) in enumerate([("Base", 0.0), ("USD +10%", -0.10), ("USD -10%", 0.10)]):
    ws.cell(r0 + 2 + j, 1, name).font = BLUE
    c = ws.cell(r0 + 2 + j, 2, shock)
    c.font, c.number_format = BLUE, PCT
rng_names = f"$A${r0 + 2}:$A${r0 + 4}"
rng_shock = f"$B${r0 + 2}:$B${r0 + 4}"
ws.cell(r0 + 5, 1, "Selected FX shock").font = BOLD
c = ws.cell(r0 + 5, 2, f"=INDEX({rng_shock},MATCH({ref['scenario'].split('!')[1]},{rng_names},0))")
c.number_format = PCT
ref["shock"] = f"Inputs!$B${r0 + 5}"
dv = DataValidation(type="list", formula1=f"={rng_names}", allow_blank=False)
ws.add_data_validation(dv)
dv.add(ref["scenario"].split("!")[1].replace("$", ""))
ws.column_dimensions["A"].width = 52
ws.column_dimensions["B"].width = 16

# ---------------- Churn Curve -------------------------------------------------
cc = wb.create_sheet("Churn Curve")
cc["A1"], cc["A1"].font = "Monthly-Plan Churn Curve and Survival", TITLE
for j, h in enumerate(["Paid tenure (months)", "Churn rate", "Survival (still active)"], start=1):
    x = cc.cell(3, j, h)
    x.font, x.fill = HDR, HFILL
for t in range(0, H + 1):
    r = 4 + t
    cc.cell(r, 1, t).font = BLACK
    if t == 0:
        cc.cell(r, 2, 0).font = BLUE
        cc.cell(r, 3, 1).font = BLUE
    else:
        cc.cell(r, 2, f"=MIN(1,{ref['haz_a']}*A{r}^{ref['haz_b']})").font = GREEN
        cc.cell(r, 3, f"=C{r - 1}*(1-B{r})").font = BLACK
    cc.cell(r, 2).number_format = PCT
    cc.cell(r, 3).number_format = PCT
SURV = f"'Churn Curve'!$C$4:$C${4 + H}"
for col, w in zip("ABC", (22, 14, 22)):
    cc.column_dimensions[col].width = w

# ---------------- Forecast ----------------------------------------------------
fs = wb.create_sheet("Forecast")
fs["A1"], fs["A1"].font = "Cohort-Based Revenue Forecast, Jan 2026 - Dec 2027 (USD)", TITLE
months = list(fc.month)
fs.cell(3, 1, "Month").font = HDR
fs.cell(3, 1).fill = HFILL
fs.cell(4, 1, "Forecast month #").font = BOLD
for j, m in enumerate(months, start=2):
    x = fs.cell(3, j, m)
    x.font, x.fill, x.alignment = HDR, HFILL, Alignment(horizontal="center")
    fs.cell(4, j, j - 1).font = BLACK
rows = {
    "trials": 5, "new_paid": 6, "new_m": 7, "new_a": 8,
    "exist_m": 10, "cohort_m": 11, "act_m": 12,
    "exist_a": 13, "cohort_a": 14, "act_a": 15, "act": 16,
    "rev_cc": 18, "rev_rep": 19,
}
labels = {
    "trials": "Trial starts", "new_paid": "New paying subscribers", "new_m": "  of which monthly plan",
    "new_a": "  of which annual plan", "exist_m": "Existing monthly subs (Dec 2025 base)",
    "cohort_m": "New monthly-plan cohorts still active", "act_m": "Active monthly-plan subs",
    "exist_a": "Existing annual subs (Dec 2025 base)", "cohort_a": "New annual-plan cohorts still active",
    "act_a": "Active annual-plan subs", "act": "Total active subscribers",
    "rev_cc": "Net revenue, constant currency ($)", "rev_rep": "Net revenue, selected FX scenario ($)",
}
for k, r in rows.items():
    fs.cell(r, 1, labels[k]).font = BOLD if k in ("act", "rev_cc", "rev_rep") else BLACK
fs.cell(9, 1, "Subscribers").font = BOLD
fs.cell(17, 1, "Revenue").font = BOLD

for j in range(2, H + 2):
    c = L(j)
    prev = L(j - 1)
    d = f"{c}$4"
    fs[f"{c}5"] = f"={ref['trials']}*(1+{ref['growth']})^({d}-1)*IF(MOD({d},12)=0,{ref['dec']},1)"
    fs[f"{c}6"] = (f"={ref['dec25']}*{ref['conv']}" if j == 2 else f"={prev}5*{ref['conv']}")
    fs[f"{c}7"] = f"={c}6*(1-{ref['ann_share']})"
    fs[f"{c}8"] = f"={c}6*{ref['ann_share']}"
    row = fc.iloc[j - 2]
    for k, val in (("exist_m", row.existing_monthly), ("exist_a", row.existing_annual)):
        x = fs.cell(rows[k], j, float(val))
        x.font = BLUE
    fs[f"{c}11"] = f"='Cohort Waterfall'!{c}{30}"
    fs[f"{c}12"] = f"={c}10+{c}11"
    fs[f"{c}14"] = f"='Cohort Waterfall'!{c}{59}"
    fs[f"{c}15"] = f"={c}13+{c}14"
    fs[f"{c}16"] = f"={c}12+{c}15"
    fs[f"{c}18"] = f"={c}12*{ref['arpu_m']}+{c}15*{ref['arpu_a']}"
    fs[f"{c}19"] = f"={c}18*(1+{ref['shock']}*{ref['non_usd']})"
    for r in range(5, 20):
        x = fs.cell(r, j)
        if x.value is None:
            continue
        x.number_format = USD if r >= 18 else NUM
        x.font = BLUE if r in (10, 13) else GREEN if r in (11, 14) else BLACK
    fs.cell(16, j).border = thin
    fs.cell(18, j).border = thin
fs.cell(10, 1).comment = Comment("Dec 2025 subscribers carried forward on the fitted churn curve by their tenure; computed in Python (run.py).", "Model")
fs.cell(13, 1).comment = Comment("Dec 2025 annual subscribers, renewing at each anniversary; computed in Python (run.py).", "Model")
fs.column_dimensions["A"].width = 40
for j in range(2, H + 2):
    fs.column_dimensions[L(j)].width = 11
fs.freeze_panes = "B5"

# ---------------- Cohort Waterfall --------------------------------------------
cw = wb.create_sheet("Cohort Waterfall")
cw["A1"], cw["A1"].font = "Cohort Waterfall: active subscribers by start month (rows) and calendar month (columns)", TITLE


def block(top, title, kind):
    cw.cell(top, 1, title).font = BOLD
    x = cw.cell(top + 1, 1, "Cohort start")
    x.font, x.fill = HDR, HFILL
    for j, m in enumerate(months, start=2):
        x = cw.cell(top + 1, j, m)
        x.font, x.fill = HDR, HFILL
    for i in range(1, H + 1):
        r = top + 1 + i
        cw.cell(r, 1, months[i - 1]).font = BLACK
        start = f"Forecast!{L(i + 1)}${8 if kind == 'a' else 7}"
        for j in range(2, H + 2):
            d = j - 1
            if d < i:
                cw.cell(r, j, 0).number_format = NUM
                continue
            if kind == "m":
                f = f"={start}*INDEX({SURV},{d - i}+1)"
            else:
                f = f"={start}*{ref['renew']}^INT(({d}-{i})/12)"
            x = cw.cell(r, j, f)
            x.font, x.number_format = GREEN, NUM
    tot = top + 2 + H
    cw.cell(tot, 1, "Total").font = BOLD
    for j in range(2, H + 2):
        x = cw.cell(tot, j, f"=SUM({L(j)}{top + 2}:{L(j)}{top + 1 + H})")
        x.font, x.number_format, x.border = BOLD, NUM, thin
    return tot


t1 = block(4, "Monthly plan: new payers x survival from the churn curve", "m")     # total row 30
t2 = block(33, "Annual plan: new payers x renewal rate at each anniversary", "a")  # total row 59
assert (t1, t2) == (30, 59)
cw.column_dimensions["A"].width = 16
for j in range(2, H + 2):
    cw.column_dimensions[L(j)].width = 9
cw.freeze_panes = "B6"

# ---------------- Summary ----------------------------------------------------
sm = wb.create_sheet("Summary", 0)
sm["A1"], sm["A1"].font = "Subscription Revenue Forecast - Summary", TITLE
sm["A2"] = "Simulated business for portfolio use; not real company data. Change inputs on the Inputs tab."
sm["A2"].font = Font(name=F, italic=True)
hdr = ["", "FY2026", "FY2027", "YoY growth"]
for j, h in enumerate(hdr, start=1):
    x = sm.cell(4, j, h)
    x.font, x.fill = HDR, HFILL
lines = [("Net revenue, constant currency ($)", 18, USD), ("Net revenue, selected FX scenario ($)", 19, USD),
         ("Ending active subscribers", 16, NUM)]
for i, (lab, r, fmt) in enumerate(lines, start=5):
    sm.cell(i, 1, lab).font = BOLD
    if r == 16:
        sm.cell(i, 2, f"=Forecast!M{r}")
        sm.cell(i, 3, f"=Forecast!Y{r}")
    else:
        sm.cell(i, 2, f"=SUM(Forecast!B{r}:M{r})")
        sm.cell(i, 3, f"=SUM(Forecast!N{r}:Y{r})")
    sm.cell(i, 4, f"=IFERROR(C{i}/B{i}-1,0)")
    for j in (2, 3):
        sm.cell(i, j).number_format, sm.cell(i, j).font = fmt, GREEN
    sm.cell(i, 4).number_format = PCT
sm.cell(8, 1, "FX impact vs constant currency ($)").font = BOLD
sm.cell(8, 2, "=B6-B5").number_format = USD
sm.cell(8, 3, "=C6-C5").number_format = USD
sm.cell(10, 1, "Check against Python model").font = BOLD
sm.cell(11, 1, "Python forecast, constant currency ($)")
for j, v in ((2, py_fy26), (3, py_fy27)):
    x = sm.cell(11, j, float(v))
    x.font, x.number_format = BLUE, USD
    x.comment = Comment("Output of subrev.analysis.forecast() in run.py.", "Model")
sm.cell(12, 1, "Difference (Excel vs Python)")
for j in (2, 3):
    cl = L(j)
    x = sm.cell(12, j, f"=IFERROR({cl}5/{cl}11-1,0)")
    x.number_format = PCT
sm.cell(14, 1, "Scenario in use").font = BOLD
sm.cell(14, 2, f"={ref['scenario']}").font = GREEN
sm.column_dimensions["A"].width = 42
for col in "BCD":
    sm.column_dimensions[col].width = 16

for sh in wb.worksheets:
    for row in sh.iter_rows():
        for x in row:
            if x.font is not None and x.font.name != F:
                x.font = Font(name=F, bold=x.font.bold, italic=x.font.italic,
                              color=x.font.color, size=x.font.size)

wb.save("subscription_forecast_excel_model.xlsx")
print("saved; python FY26 %.0f FY27 %.0f" % (py_fy26, py_fy27))
